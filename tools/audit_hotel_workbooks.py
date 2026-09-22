#!/usr/bin/env python3
"""
JoTrip private hotel workbook audit helper.

This script never uploads data. It reads Excel files locally and writes a
private audit JSON under ./private-data/, which is gitignored.

Usage:
  python3 -m pip install openpyxl
  python3 tools/audit_hotel_workbooks.py /path/to/file.xlsx
"""

from __future__ import annotations

import json
import re
import sys
import uuid
from pathlib import Path
from typing import Any

import openpyxl

ALL_MARKET_RE = re.compile(r"\ball\s*markets?\b|tất\s*cả\s*(?:các\s*)?thị\s*trường", re.I)
PERCENT_RE = re.compile(r"(?:giảm|discount)\s*(\d+(?:[.,]\d+)?)\s*%", re.I)
DATE_RE = re.compile(
    r"\b(\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?)\b"
)


def text_of(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def audit_sheet(ws) -> dict[str, Any]:
    markers = []
    percentages = []
    date_tokens = []
    numeric_cells = []

    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            text = text_of(value)
            if not text:
                continue

            if ALL_MARKET_RE.search(text):
                markers.append({
                    "type": "ALL_MARKET",
                    "cell": cell.coordinate,
                    "text": text[:500],
                })

            for m in PERCENT_RE.finditer(text):
                percentages.append({
                    "cell": cell.coordinate,
                    "percent": float(m.group(1).replace(",", ".")),
                    "text": text[:500],
                })

            for d in DATE_RE.findall(text):
                date_tokens.append({
                    "cell": cell.coordinate,
                    "token": d,
                })

            if isinstance(value, (int, float)) and value >= 100_000:
                numeric_cells.append({
                    "cell": cell.coordinate,
                    "value": value,
                })

    state = "REVIEW"
    reasons = []

    if markers:
        state = "CANDIDATE"
    else:
        reasons.append("no_clear_all_market_marker")

    if len(percentages) > 1:
        reasons.append("multiple_discount_candidates_review_required")

    return {
        "sheet": ws.title,
        "state": state,
        "reasons": reasons,
        "all_market_markers": markers,
        "discount_candidates": percentages,
        "date_tokens": date_tokens[:100],
        "numeric_rate_candidates": numeric_cells[:500],
    }


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Pass one or more .xlsx file paths")

    out_dir = Path("private-data")
    out_dir.mkdir(parents=True, exist_ok=True)

    report = {
        "audit_id": str(uuid.uuid4()),
        "files": [],
    }

    for raw_path in sys.argv[1:]:
        path = Path(raw_path).expanduser()
        wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        pq_sheets = [
            name for name in wb.sheetnames
            if name.strip().upper().startswith("PQ")
            or "PHU QUOC" in name.upper()
            or "PHÚ QUỐC" in name.upper()
        ]

        report["files"].append({
            "source_name": path.name,
            "sheet_count": len(wb.sheetnames),
            "candidate_sheet_count": len(pq_sheets),
            "sheets": [audit_sheet(wb[name]) for name in pq_sheets],
        })

    output = out_dir / "hotel-workbook-audit.json"
    output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"Wrote private audit to {output}")
    print("Do not commit this file.")


if __name__ == "__main__":
    main()
